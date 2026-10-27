from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def literal(cell, value):
    # Force text even if customer/product starts with Excel formula characters.
    cell.value = str(value)
    cell.data_type = 's'


def export_review(row, path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Review'
    metadata = [('Status', 'APPROVED' if row['approved_hash'] == row['hash'] else 'DRAFT - REVIEW REQUIRED'),
                ('Order ID', row['id']), ('Revision', row['revision']), ('Review SHA256', row['hash']),
                ('Customer', row['record']['customer']), ('Date', row['record']['sale_date']),
                ('Order reference', row['record']['order_ref']), ('Source SHA256', row['record']['source_hash']),
                ('Price status', row['totals']['price_status'])]
    for i, (label, value) in enumerate(metadata, 1):
        literal(ws.cell(i, 1), label)
        literal(ws.cell(i, 2), value)
    headers = ['Line', 'Product', 'Unit', 'Quantity', 'Unit price VND', 'Amount VND', 'Price basis']
    ws.append([])
    ws.append(headers)
    for item in row['record']['lines']:
        r = ws.max_row + 1
        for c, key in [(1, 'line_id'), (2, 'product'), (3, 'unit'), (7, 'price_basis')]:
            literal(ws.cell(r, c), item[key])
        ws.cell(r, 4, float(Decimal(item['quantity'])))
        if item['unit_price_vnd'] is not None:
            ws.cell(r, 5, float(Decimal(item['unit_price_vnd'])))
            ws.cell(r, 6, float(Decimal(item['quantity']) * Decimal(item['unit_price_vnd'])))
        else:
            literal(ws.cell(r, 5), 'PENDING')
    ws.append(['Known subtotal VND', int(row['totals']['known_total_vnd'])])
    ws.append(['Final total VND', None if row['totals']['final_total_vnd'] is None else int(row['totals']['final_total_vnd'])])
    for cell in ws[11]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='164E63')
    for col, width in {'A': 24, 'B': 72, 'C': 14, 'D': 16, 'E': 22, 'F': 22, 'G': 45}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = 'D12'
    ws.auto_filter.ref = f'A11:G{11 + len(row["record"]["lines"])}'
    routing = row['record'].get('routing')
    if routing:
        literal(ws.cell(1, 4), 'Routing status')
        literal(ws.cell(1, 5), routing['status'])
        literal(ws.cell(2, 4), 'Destination')
        literal(ws.cell(2, 5), routing['destination'] or 'PENDING')
        literal(ws.cell(3, 4), 'Routing rules hash')
        literal(ws.cell(3, 5), routing['rules_hash'])
    wb.save(path)
