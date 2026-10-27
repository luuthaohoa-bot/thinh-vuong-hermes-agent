import copy
import io
import json
import tempfile
import unittest
import shutil
import uuid
from unittest.mock import MagicMock, patch
from pathlib import Path
from pypdf import PdfWriter
from openpyxl import load_workbook
from hermes.store import Store
from hermes.intake import intake, SidecarOCR
from hermes.export import export_review
from hermes.telegram import parse_update
from hermes.model import validate

BASE = Path(__file__).resolve().parents[1]


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        # Normal inherited ACLs also work inside Windows restricted runners.
        self.root = Path(tempfile.gettempdir()) / ('hermes-test-' + uuid.uuid4().hex)
        self.root.mkdir()
        self.store = Store(self.root / 'db.sqlite')
        self.record = json.loads((BASE / 'examples/order.json').read_text())
        self.record['routing'] = {'status': 'assigned', 'destination': 'TEST ONLY', 'rules_hash': 'synthetic'}

    def tearDown(self):
        self.store.close()
        shutil.rmtree(self.root)

    def test_approval_duplicate_and_retry(self):
        row = self.store.create(self.record)
        self.assertEqual(row['id'], self.store.create(self.record)['id'])
        with self.assertRaises(ValueError): self.store.publish_mock(row['id'])
        with self.assertRaises(ValueError): self.store.approve(row['id'], 'stale', 'owner')
        self.store.approve(row['id'], row['hash'], 'owner')
        self.assertTrue(self.store.publish_mock(row['id'])['inserted'])
        self.assertFalse(self.store.publish_mock(row['id'])['inserted'])
        self.assertIsNone(row['totals']['final_total_vnd'])

    def test_settlement_requires_new_review(self):
        row = self.store.create(self.record)
        self.store.approve(row['id'], row['hash'], 'owner')
        self.store.publish_mock(row['id'])
        changed = copy.deepcopy(row['record'])
        changed['lines'][0].update(unit_price_vnd='130000', price_basis='Market 2026-10-25')
        updated = self.store.revise(row['id'], changed, row['hash'])
        self.assertEqual(updated['revision'], 2)
        self.assertEqual(updated['totals']['final_total_vnd'], '1542000')
        with self.assertRaises(ValueError): self.store.publish_mock(row['id'])
        with self.assertRaises(ValueError): self.store.approve(row['id'], row['hash'], 'owner')
        self.store.approve(row['id'], updated['hash'], 'owner')
        self.assertTrue(self.store.publish_mock(row['id'])['inserted'])
        self.assertEqual(self.store.db.execute('SELECT count(*) FROM mock_writes').fetchone()[0], 2)

    def test_business_duplicate_with_new_pdf(self):
        self.store.create(self.record)
        self.record['source_hash'] = 'b' * 64
        with self.assertRaises(ValueError): self.store.create(self.record)

    def test_noop_edit_preserves_approval_and_stale_edit_fails(self):
        row = self.store.create(self.record)
        self.store.approve(row['id'], row['hash'], 'owner')
        row = self.store.revise(row['id'], row['record'], row['hash'])
        self.assertEqual(row['approved_hash'], row['hash'])
        with self.assertRaises(ValueError): self.store.revise(row['id'], row['record'], 'old')

    def test_invalid_quantities(self):
        for q in ['NaN', 'Infinity', '-1', '0', '1,2']:
            self.record['lines'][0]['quantity'] = q
            with self.assertRaises(ValueError): validate(self.record)

    def test_excel_injection_and_pending(self):
        self.record['customer'] = '=HYPERLINK("evil")'
        self.record['lines'][0]['product'] = '=1+1'
        row = self.store.create(self.record)
        file = self.root / 'review.xlsx'
        export_review(row, file)
        wb = load_workbook(file)
        ws = wb.active
        self.assertEqual(ws['B5'].data_type, 's')
        self.assertEqual(ws['B12'].data_type, 's')
        self.assertEqual(ws['E12'].value, 'PENDING')
        self.assertIsNone(ws['B15'].value)
        wb.close()

    def test_scanned_pdf_needs_ocr(self):
        writer = PdfWriter(); writer.add_blank_page(width=300, height=400)
        path = self.root / 'scan.pdf'
        with path.open('wb') as f: writer.write(f)
        result = intake(path)
        self.assertEqual(result['pages'][0]['method'], 'needs_ocr')
        self.assertEqual(result['sales_records'], [])
        result = intake(path, SidecarOCR({'1': 'manual transcription'}))
        self.assertEqual(result['pages'][0]['method'], 'ocr_unverified')

    def test_telegram_guard(self):
        update = {'update_id': 1, 'message': {'chat': {'id': 123, 'type': 'private'},
                  'document': {'file_id': 'abc', 'mime_type': 'application/pdf', 'file_size': 100}}}
        self.assertEqual(parse_update(update, 'secret', 'secret', [123])['action'], 'queue_for_review_only')
        for secret, chats in [('bad', [123]), ('secret', [456])]:
            with self.assertRaises(ValueError): parse_update(update, secret, 'secret', chats)
        with self.assertRaises(ValueError): parse_update(update, '', '', [123])

    def test_bhx_helpers(self):
        from hermes.bhx_reader import number, clean, pdf_members
        self.assertEqual(number('1,234.5'), 1234.5)
        self.assertEqual(clean(' a  b '), 'a b')
        with self.assertRaises(ValueError): number('NaN')
        with self.assertRaises(ValueError): pdf_members(b'anything', 'x.txt')

    def test_bhx_table_to_pending_candidates_and_fail_closed(self):
        from hermes.bhx import candidates
        path = self.root / 'allocation.pdf'
        path.write_bytes(b'%PDF-1.4\nfixture')
        page = MagicMock()
        page.extract_text.return_value = 'BIÊN BẢN CHIA HÀNG Mã - Tên kho: 123 - BHX_HCM_Q1 - Kho A Nhà cung cấp:'
        page.extract_tables.return_value = [[['1', 'BHX_STORE_1', '', '1', 'CÁ THU', '8.4', 'Kg']]]
        document = MagicMock()
        document.pages = [page]
        with patch('hermes.bhx_reader.pdfplumber.open') as opener:
            opener.return_value.__enter__.return_value = document
            result = candidates(path, '2026-10-24')
            self.assertEqual(result['records'][0]['lines'][0]['quantity'], '8.4')
            self.assertIsNone(result['records'][0]['lines'][0]['unit_price_vnd'])
            page.extract_tables.return_value = []
            with self.assertRaises(ValueError): candidates(path, '2026-10-24')

    def test_approval_and_dedupe_survive_restart(self):
        row = self.store.create(self.record)
        self.store.approve(row['id'], row['hash'], 'owner')
        self.store.publish_mock(row['id'])
        other = Store(self.root / 'db.sqlite')
        try:
            self.assertFalse(other.publish_mock(row['id'])['inserted'])
            changed = other.get(row['id'])['record']
            changed['lines'][0]['quantity'] = '9'
            other.revise(row['id'], changed, row['hash'])
            with self.assertRaises(ValueError): self.store.publish_mock(row['id'])
        finally:
            other.close()


if __name__ == '__main__':
    unittest.main()
